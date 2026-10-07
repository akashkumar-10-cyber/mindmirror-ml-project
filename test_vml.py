import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
import docx2pdf

def run_test():
    doc = docx.Document()
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    header = section.header
    p_head = header.paragraphs[0]
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER

    wm_path = r"C:\Users\Dell\.gemini\antigravity\scratch\ML project\siragu_watermark_final.png"
    r_dummy = p_head.add_run()
    pic = r_dummy.add_picture(wm_path, width=Inches(4.5))

    blip = r_dummy._r.xpath('.//a:blip')[0]
    r_id = blip.get(qn('r:embed'))
    print('Image relationship ID in header:', r_id)

    p_head._p.remove(r_dummy._r)

    vml_xml = (
        f'<w:r {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">\n'
        f'  <w:pict>\n'
        f'    <v:shape id="WordPictureWatermark1" style="position:absolute;margin-left:0;margin-top:0;width:380pt;height:90pt;z-index:-251658240;mso-position-horizontal:center;mso-position-horizontal-relative:margin;mso-position-vertical:center;mso-position-vertical-relative:margin" fillcolor="none" stroked="f">\n'
        f'      <v:imagedata r:id="{r_id}" o:title="SIRAGU"/>\n'
        f'    </v:shape>\n'
        f'  </w:pict>\n'
        f'</w:r>'
    )
    p_head._p.append(parse_xml(vml_xml))

    doc.add_paragraph('PAGE 1: TITLE PAGE (NO WATERMARK)')
    doc.add_page_break()
    doc.add_paragraph('PAGE 2: WITH WATERMARK IN BACKGROUND')
    doc.add_page_break()
    doc.add_paragraph('PAGE 3: WITH WATERMARK IN BACKGROUND')

    test_docx = r"C:\Users\Dell\.gemini\antigravity\scratch\ML project\test_vml_wm.docx"
    test_pdf = r"C:\Users\Dell\.gemini\antigravity\scratch\ML project\test_vml_wm.pdf"
    doc.save(test_docx)
    print("Saved test VML docx")

    docx2pdf.convert(test_docx, test_pdf)
    print("VML Watermark converted perfectly to PDF!")

if __name__ == "__main__":
    run_test()
