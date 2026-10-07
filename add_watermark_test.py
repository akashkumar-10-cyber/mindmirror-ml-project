import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_test():
    doc = docx.Document()
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    header = section.header
    p_head = header.paragraphs[0]

    wm_path = r"C:\Users\Dell\.gemini\antigravity\scratch\ML project\siragu_watermark_final.png"
    r = p_head.add_run()
    r.add_picture(wm_path, width=Inches(5.0))

    drawing = r._r.xpath('.//w:drawing')[0]
    inline = drawing.xpath('.//wp:inline')[0]

    anchor_xml = (
        f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="251658240" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">\n'
        f'  <wp:simplePos x="0" y="0"/>\n'
        f'  <wp:positionH relativeFrom="page">\n'
        f'    <wp:align>center</wp:align>\n'
        f'  </wp:positionH>\n'
        f'  <wp:positionV relativeFrom="page">\n'
        f'    <wp:align>center</wp:align>\n'
        f'  </wp:positionV>\n'
        f'</wp:anchor>'
    )
    anchor_elem = parse_xml(anchor_xml)
    for child in list(inline):
        anchor_elem.append(child)

    drawing.remove(inline)
    drawing.append(anchor_elem)

    # Page 1 (Title)
    doc.add_paragraph('PAGE 1: TITLE PAGE (NO WATERMARK)')
    doc.add_page_break()
    # Page 2
    doc.add_paragraph('PAGE 2: WITH WATERMARK IN BACKGROUND')
    doc.add_page_break()
    # Page 3
    doc.add_paragraph('PAGE 3: WITH WATERMARK IN BACKGROUND')

    out_p = r"C:\Users\Dell\.gemini\antigravity\scratch\ML project\test_watermark.docx"
    doc.save(out_p)
    print("Saved test watermark docx successfully!")

if __name__ == "__main__":
    create_test()
