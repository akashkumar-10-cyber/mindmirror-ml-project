import docx

doc = docx.Document(r"C:\Users\Dell\Downloads\MindMirror_AI_PBL_Report_Updated_Final.docx")
abs_idx = [i for i, p in enumerate(doc.paragraphs) if "ABSTRACT" in p.text][0]
print("Abstract paragraph index:", abs_idx)

for i, el in enumerate(doc._body._element):
    tag = el.tag.split("}")[-1]
    if tag == "p":
        p = docx.text.paragraph.Paragraph(el, doc)
        draws = len(p._p.xpath(".//w:drawing"))
        print(f"{i:02d} [P] (draw={draws}): {repr(p.text[:60])}")
        if "ABSTRACT" in p.text:
            break
    elif tag == "tbl":
        t = docx.table.Table(el, doc)
        print(f"{i:02d} [TBL {len(t.rows)}x{len(t.columns)}]: {repr(t.cell(0,0).text[:60])}")
